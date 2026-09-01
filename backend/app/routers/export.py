"""Excel 导出路由：按原样导出各表为 .xlsx。"""
import io
from urllib.parse import quote
from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from app.database import get_db
from app.models import (
    EmployeeTitle, TalentAccount, SubsidyPolicy, SubsidyApplication, SubsidyPayment, SysUser
)
from app.security import get_current_user
from app.encryption import decrypt_text
import openpyxl
from openpyxl.styles import Font

router = APIRouter(prefix="/api/export", tags=["export"])


def _xlsx_response(wb: openpyxl.Workbook, filename: str):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    headers = {
        "Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"
    }
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )


def _header_row(ws, headers):
    for i, h in enumerate(headers, 1):
        c = ws.cell(1, i, h)
        c.font = Font(bold=True)


@router.get("/titles")
def export_titles(db: Session = Depends(get_db), current_user: SysUser = Depends(get_current_user)):
    """导出职称表（与 职称汇总表_证书补全.xlsx 同结构）。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    _header_row(ws, ["序号", "姓名", "学历", "职称类型", "职称认定时间", "专业", "学科", "下一阶段申请", "备注"])
    rows = db.query(EmployeeTitle).filter(EmployeeTitle.is_deleted == False)\
        .order_by(EmployeeTitle.seq_no.asc(), EmployeeTitle.id.asc()).all()
    for i, t in enumerate(rows, 1):
        ws.cell(i + 1, 1, t.seq_no or i)
        ws.cell(i + 1, 2, t.name)
        ws.cell(i + 1, 3, t.education)
        ws.cell(i + 1, 4, t.title_name)
        ws.cell(i + 1, 5, t.title_date.strftime("%Y.%m.%d") if t.title_date else "")
        ws.cell(i + 1, 6, t.major_field)
        ws.cell(i + 1, 7, t.discipline)
        ws.cell(i + 1, 8, t.next_stage_apply)
        ws.cell(i + 1, 9, t.remark)
    return _xlsx_response(wb, "职称汇总表_导出.xlsx")


@router.get("/accounts")
def export_accounts(db: Session = Depends(get_db), current_user: SysUser = Depends(get_current_user)):
    """导出人才账号表（与 个人数币账号总明细.xlsx 同结构，账号解密明文）。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    _header_row(ws, ["序号", "姓名", "数币银行", "数币账号", "普通账户", "普通账号"])
    rows = db.query(TalentAccount).filter(TalentAccount.is_deleted == False).order_by(TalentAccount.id.asc()).all()
    for i, a in enumerate(rows, 1):
        ws.cell(i + 1, 1, i)
        ws.cell(i + 1, 2, a.name)
        ws.cell(i + 1, 3, a.digital_bank)
        ws.cell(i + 1, 4, decrypt_text(a.digital_account_encrypted))
        ws.cell(i + 1, 5, a.bank_name)
        ws.cell(i + 1, 6, decrypt_text(a.bank_account_encrypted))
    return _xlsx_response(wb, "个人数币账号总明细_导出.xlsx")


@router.get("/policies")
def export_policies(db: Session = Depends(get_db), current_user: SysUser = Depends(get_current_user)):
    """导出政策清单。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "政策清单"
    _header_row(ws, ["序号", "政策名称", "类别", "区域", "级别", "状态", "备注"])
    rows = db.query(SubsidyPolicy).filter(SubsidyPolicy.is_deleted == False).order_by(SubsidyPolicy.id.asc()).all()
    for i, p in enumerate(rows, 1):
        ws.cell(i + 1, 1, i)
        ws.cell(i + 1, 2, p.policy_name)
        ws.cell(i + 1, 3, p.policy_category)
        ws.cell(i + 1, 4, p.region)
        ws.cell(i + 1, 5, p.region_level)
        ws.cell(i + 1, 6, "在发" if p.status == "active" else "已结束")
        ws.cell(i + 1, 7, p.remark)
    return _xlsx_response(wb, "人才补贴政策清单_导出.xlsx")


@router.get("/subsidy")
def export_subsidy(db: Session = Depends(get_db), current_user: SysUser = Depends(get_current_user)):
    """导出补贴明细（申领 + 发放纵向展开）。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "补贴明细"
    _header_row(ws, ["序号", "政策", "姓名", "BU", "部门", "期次", "应发日期", "实发日期", "金额", "状态", "备注"])
    apps = db.query(SubsidyApplication).filter(SubsidyApplication.is_deleted == False).order_by(SubsidyApplication.id.asc()).all()
    r = 2
    for a in apps:
        policy = db.query(SubsidyPolicy).filter(SubsidyPolicy.id == a.policy_id).first()
        payments = db.query(SubsidyPayment).filter(SubsidyPayment.application_id == a.id)\
            .order_by(SubsidyPayment.payment_index.asc()).all()
        if not payments:
            ws.cell(r, 1, a.id)
            ws.cell(r, 2, policy.policy_name if policy else "")
            ws.cell(r, 3, a.name)
            ws.cell(r, 4, a.bu)
            ws.cell(r, 5, a.department)
            ws.cell(r, 10, a.total_received)
            ws.cell(r, 11, a.remark)
            r += 1
        for p in payments:
            ws.cell(r, 1, a.id)
            ws.cell(r, 2, policy.policy_name if policy else "")
            ws.cell(r, 3, a.name)
            ws.cell(r, 4, a.bu)
            ws.cell(r, 5, a.department)
            ws.cell(r, 6, p.payment_index)
            ws.cell(r, 7, p.expected_date.strftime("%Y.%m.%d") if p.expected_date else "")
            ws.cell(r, 8, p.actual_date.strftime("%Y.%m.%d") if p.actual_date else "")
            ws.cell(r, 9, p.amount)
            ws.cell(r, 10, p.status)
            ws.cell(r, 11, p.remark)
            r += 1
    return _xlsx_response(wb, "个人补贴总表_导出.xlsx")
