"""Excel 一键录入：上传与导出同结构的 xlsx，按表头动态解析并写入。"""
import shutil
import uuid
from datetime import date, datetime
from pathlib import Path

from fastapi import APIRouter, Depends, File, UploadFile
from sqlalchemy.orm import Session

from app.database import get_db
from app.models import (
    EmployeeTitle, TalentAccount, SubsidyApplication, SubsidyPolicy,
    SysUser, AuditLog, SubsidyPayment,
)
from app.security import get_current_user
from app.encryption import encrypt_text, mask_account
from app import config

router = APIRouter(prefix="/api/import", tags=["import"])

IMPORT_DIR = config.UPLOAD_DIR / "import"
IMPORT_DIR.mkdir(parents=True, exist_ok=True)


def _parse_date(v):
    if not v:
        return None
    if isinstance(v, (int, float)):
        try:
            return openpyxl_dt(v)
        except Exception:
            return None
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y年%m月%d日", "%Y-%m", "%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def openpyxl_dt(v):
    import openpyxl.utils
    return openpyxl.utils.datetime.from_excel(v).date()


def _header_map(ws):
    """返回 表头中文 -> 列号(1-based)。"""
    headers = {}
    for c in range(1, ws.max_column + 1):
        h = ws.cell(1, c).value
        if h:
            headers[str(h).strip()] = c
    return headers


def _msg(ok, msg):
    return {"ok": ok, "message": msg}


@router.post("/titles")
def import_titles(db: Session = Depends(get_db), current_user: SysUser = Depends(get_current_user),
                 file: UploadFile = File(...)):
    import openpyxl
    saved = IMPORT_DIR / f"{uuid.uuid4().hex}_{file.filename}"
    with saved.open("wb") as out:
        shutil.copyfileobj(file.file, out)
    wb = openpyxl.load_workbook(str(saved), data_only=True)
    ws = wb.active
    h = _header_map(ws)
    need = ["姓名", "职称类型"]
    if not all(k in h for k in need):
        return _msg(False, "未找到必需的表头（需含：姓名、职称类型）")
    series_map = {"工程类": "engineering", "技能类": "technician"}
    level_map = {"高级工程师": "sub_senior", "工程师": "middle", "助理工程师": "junior",
                 "一级（高级技师）": "level_1", "二级（技师）": "level_2", "三级（高级工）": "level_3",
                 "助理研究员": "junior_researcher"}
    n = 0
    for ridx in range(2, ws.max_row + 1):
        name = ws.cell(ridx, h["姓名"]).value
        title_name = ws.cell(ridx, h["职称类型"]).value
        if not name or not str(name).strip() or not title_name:
            continue
        title_name = str(title_name).strip()
        existing = db.query(EmployeeTitle).filter(
            EmployeeTitle.name == str(name).strip(),
            EmployeeTitle.title_name == title_name,
            EmployeeTitle.is_deleted == False).first()
        if existing:
            continue
        seq = ws.cell(ridx, h.get("序号", 1)).value if "序号" in h else None
        t = EmployeeTitle(
            seq_no=int(seq) if str(seq).isdigit() else None,
            name=str(name).strip(),
            education=ws.cell(ridx, h["学历"]).value if "学历" in h else None,
            title_name=title_name,
            title_series=series_map.get(ws.cell(ridx, h["职称系列"]).value if "职称系列" in h else "", "engineering"),
            title_level=level_map.get(title_name, "middle"),
            title_date=_parse_date(ws.cell(ridx, h["职称认定时间"]).value) if "职称认定时间" in h else None,
            major_field=ws.cell(ridx, h["专业"]).value if "专业" in h else None,
            discipline=ws.cell(ridx, h["学科"]).value if "学科" in h else None,
            next_stage_apply=ws.cell(ridx, h["下一阶段申请"]).value if "下一阶段申请" in h else None,
            remark=ws.cell(ridx, h["备注"]).value if "备注" in h else None,
            created_by=current_user.id,
        )
        db.add(t)
        n += 1
    db.commit()
    _log(db, current_user, "import", "employee_title", 0, after={"added": n})
    return _msg(True, f"成功导入 {n} 条职称记录（已存在的自动跳过）")


@router.post("/accounts")
def import_accounts(db: Session = Depends(get_db), current_user: SysUser = Depends(get_current_user),
                   file: UploadFile = File(...)):
    import openpyxl
    saved = IMPORT_DIR / f"{uuid.uuid4().hex}_{file.filename}"
    with saved.open("wb") as out:
        shutil.copyfileobj(file.file, out)
    wb = openpyxl.load_workbook(str(saved), data_only=True)
    ws = wb.active
    h = _header_map(ws)
    if "姓名" not in h:
        return _msg(False, "未找到必需的表头（需含：姓名）")
    n = 0
    for ridx in range(2, ws.max_row + 1):
        name = ws.cell(ridx, h["姓名"]).value
        if not name or not str(name).strip():
            continue
        d_acc = ws.cell(ridx, h["数币账号"]).value if "数币账号" in h else None
        b_acc = ws.cell(ridx, h["普通账号"]).value if "普通账号" in h else None
        existing = db.query(TalentAccount).filter(
            TalentAccount.name == str(name).strip(), TalentAccount.is_deleted == False).first()
        if existing:
            # 仅补全缺失的账号
            if not existing.digital_account_encrypted and d_acc:
                existing.digital_account_encrypted = encrypt_text(str(d_acc))
                existing.digital_account_mask = mask_account(str(d_acc))
            if not existing.bank_account_encrypted and b_acc:
                existing.bank_account_encrypted = encrypt_text(str(b_acc))
                existing.bank_account_mask = mask_account(str(b_acc))
            continue
        a = TalentAccount(
            name=str(name).strip(),
            digital_bank=ws.cell(ridx, h["数币银行"]).value if "数币银行" in h else None,
            bank_name=ws.cell(ridx, h["普通账户"]).value if "普通账户" in h else None,
            created_by=current_user.id,
        )
        if d_acc:
            a.digital_account_encrypted = encrypt_text(str(d_acc))
            a.digital_account_mask = mask_account(str(d_acc))
        if b_acc:
            a.bank_account_encrypted = encrypt_text(str(b_acc))
            a.bank_account_mask = mask_account(str(b_acc))
        db.add(a)
        n += 1
    db.commit()
    _log(db, current_user, "import", "talent_account", 0, after={"added": n})
    return _msg(True, f"成功导入 {n} 条人才账号（已存在的自动跳过/补全）")


@router.post("/subsidy")
def import_subsidy(db: Session = Depends(get_db), current_user: SysUser = Depends(get_current_user),
                  file: UploadFile = File(...)):
    """导入补贴明细（申领 + 发放纵向展开，与导出同结构）。"""
    import openpyxl
    saved = IMPORT_DIR / f"{uuid.uuid4().hex}_{file.filename}"
    with saved.open("wb") as out:
        shutil.copyfileobj(file.file, out)
    wb = openpyxl.load_workbook(str(saved), data_only=True)
    ws = wb.active
    h = _header_map(ws)
    if "姓名" not in h or "政策" not in h:
        return _msg(False, "未找到必需的表头（需含：姓名、政策）")
    policy_map = {p.policy_name: p for p in db.query(SubsidyPolicy).all()}
    app_cache = {}
    n_app, n_pay = 0, 0
    for ridx in range(2, ws.max_row + 1):
        name = ws.cell(ridx, h["姓名"]).value
        policy_name = ws.cell(ridx, h["政策"]).value
        if not name or not str(name).strip() or not policy_name:
            continue
        name = str(name).strip()
        policy_name = str(policy_name).strip()
        policy = policy_map.get(policy_name)
        key = (name, policy_name)
        app = app_cache.get(key)
        if not app:
            app = db.query(SubsidyApplication).filter(
                SubsidyApplication.name == name,
                SubsidyApplication.policy_id == (policy.id if policy else None),
                SubsidyApplication.is_deleted == False).first()
            if not app:
                app = SubsidyApplication(
                    name=name, policy_id=policy.id if policy else None,
                    bu=ws.cell(ridx, h["BU"]).value if "BU" in h else None,
                    department=ws.cell(ridx, h["部门"]).value if "部门" in h else None,
                    total_expected=0, total_received=0, status="ongoing",
                    remark=ws.cell(ridx, h["备注"]).value if "备注" in h else None,
                    created_by=current_user.id,
                )
                db.add(app)
                db.flush()
                n_app += 1
            app_cache[key] = app
        # 发放行（有期次/金额）
        idx = ws.cell(ridx, h["期次"]).value if "期次" in h else None
        amount = ws.cell(ridx, h["金额"]).value if "金额" in h else None
        if idx is not None and amount is not None:
            idx = int(idx) if str(idx).strip().isdigit() else 0
            amt = float(amount) if str(amount).strip() not in ("", "None") else 0
            exist = db.query(SubsidyPayment).filter(
                SubsidyPayment.application_id == app.id,
                SubsidyPayment.payment_index == idx).first()
            if not exist:
                st = (ws.cell(ridx, h["状态"]).value or "pending")
                st = "pending" if st not in ("paid", "confirmed", "completed", "pending_confirm", "rejected") else st
                p = SubsidyPayment(
                    application_id=app.id, payment_index=idx,
                    expected_date=_parse_date(ws.cell(ridx, h["应发日期"]).value) if "应发日期" in h else None,
                    actual_date=_parse_date(ws.cell(ridx, h["实发日期"]).value) if "实发日期" in h else None,
                    amount=amt, status=st,
                    remark=ws.cell(ridx, h["备注"]).value if "备注" in h else None,
                    created_by=current_user.id,
                )
                db.add(p)
                n_pay += 1
    db.commit()
    # 重算累计到账
    for app in app_cache.values():
        total = db.query(SubsidyPayment).filter(
            SubsidyPayment.application_id == app.id,
            SubsidyPayment.status.in_(("paid", "confirmed", "completed"))).with_entities(SubsidyPayment.amount).all()
        app.total_received = round(sum((x[0] or 0) for x in total), 2)
    db.commit()
    _log(db, current_user, "import", "subsidy_application", 0, after={"apps": n_app, "payments": n_pay})
    return _msg(True, f"成功导入 {n_app} 条申领、{n_pay} 笔发放")


def _log(db, user, action, btype, bid, before=None, after=None):
    import json
    db.add(AuditLog(
        user_id=user.id, user_name=user.display_name, action=action,
        business_type=btype, business_id=bid,
        after_value=json.dumps(after, ensure_ascii=False, default=str) if after else None))
    db.commit()
